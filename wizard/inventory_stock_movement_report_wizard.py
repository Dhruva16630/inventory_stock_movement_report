from datetime import datetime, time
from io import BytesIO
import base64

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from odoo import models, fields, api
from odoo.exceptions import ValidationError


class InventoryStockMovementReportWizard(models.TransientModel):
    _name = "inventory.stock.movement.report.wizard"
    _description = "Inventory Stock Movement Report Wizard"

    from_date = fields.Date(
        string="From Date",
        required=True,
    )

    to_date = fields.Date(
        string="To Date",
        required=True,
    )

    excel_file = fields.Binary(
        string="Excel File",
        readonly=True,
    )

    excel_filename = fields.Char(
        string="Excel Filename",
        readonly=True,
    )

    @api.constrains("from_date", "to_date")
    def _check_dates(self):
        for record in self:
            if record.from_date > record.to_date:
                raise ValidationError(
                    "From Date cannot be later than To Date."
                )

    def _get_date_range(self):
        """Return datetime boundaries for the selected dates."""

        self.ensure_one()

        from_datetime = datetime.combine(
            self.from_date,
            time.min,
        )

        to_datetime = datetime.combine(
            self.to_date,
            time.max,
        )

        return from_datetime, to_datetime

    def action_generate_report(self):
        self.ensure_one()

        from_datetime, to_datetime = self._get_date_range()

        return {
            "type": "ir.actions.act_window",
            "name": "Stock Movement Report",
            "res_model": "inventory.stock.movement.report",
            "view_mode": "list,graph,pivot",
            "domain": [
                ("scheduled_date", ">=", from_datetime),
                ("scheduled_date", "<=", to_datetime),
            ],
        }

    def action_export_excel(self):
        self.ensure_one()

        from_datetime, to_datetime = self._get_date_range()
        
        reports = self.env[
            "inventory.stock.movement.report"
        ].search([
            ("scheduled_date", ">=", from_datetime),
            ("scheduled_date", "<=", to_datetime),
        ])
        
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Stock Movement"

        headers = [
            "Product",
            "Product Category",
            "Quantity",
            "Source Location",
            "Destination Location",
            "Reference",
            "Scheduled Date",
            "Status",
            "Picking Type",
            "Warehouse",
        ]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(
                row=1,
                column=column,
                value=header,
            )

            cell.font = Font(
                bold=True,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        row = 2
        for report in reports:
            sheet.cell(
                row=row,
                column=1,
                value=report.product.display_name,
            )

            sheet.cell(
                row=row,
                column=2,
                value=report.product_category.display_name,
            )

            sheet.cell(
                row=row,
                column=3,
                value=report.qty,
            )

            sheet.cell(
                row=row,
                column=4,
                value=report.source_location.display_name,
            )

            sheet.cell(
                row=row,
                column=5,
                value=report.destination_location.display_name,
            )

            sheet.cell(
                row=row,
                column=6,
                value=report.reference,
            )

            scheduled_date_cell = sheet.cell(
                row=row,
                column=7,
                value=report.scheduled_date,
            )

            scheduled_date_cell.number_format = "yyyy-mm-dd hh:mm:ss"

            sheet.cell(
                row=row,
                column=8,
                value=report.state,
            )

            sheet.cell(
                row=row,
                column=9,
                value=report.picking_type.display_name,
            )

            sheet.cell(
                row=row,
                column=10,
                value=report.warehouse.display_name,
            )

            row += 1

        for column_cells in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    cell_length = len(
                        str(cell.value)
                    )

                    if cell_length > max_length:
                        max_length = cell_length

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                40,
            )

        sheet.freeze_panes = "A2"

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        excel_file = base64.b64encode(
            output.read()
        )

        self.write({
            "excel_file": excel_file,
            "excel_filename": "stock_movement_report.xlsx",
        })

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": self.id,
            "target": "new",
        }